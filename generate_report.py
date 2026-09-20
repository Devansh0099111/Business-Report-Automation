import pandas as pd

# Load Excel data
file_path = "input_data/Business_Data.xlsx"

df = pd.read_excel(file_path)

print("Business data loaded successfully!")
print()

# Calculate total sales
total_sales = df["Sales"].sum()

# Calculate total expenses
total_expenses = df["Expenses"].sum()

# Calculate total profit
total_profit = total_sales - total_expenses

# Calculate profit margin
profit_margin = total_profit / total_sales

# Calculate total orders
total_orders = len(df)

# Sales by product
product_sales = df.groupby("Product")["Sales"].sum()

# Best-performing product
top_product = product_sales.idxmax()

# Sales by region
region_sales = df.groupby("Region")["Sales"].sum()

# Best-performing region
top_region = region_sales.idxmax()

# Display results
print("Top Product:", top_product)
print("Top Region:", top_region)

print()
print("----- Sales by Product -----")
print(product_sales)

print()
print("----- Sales by Region -----")
print(region_sales)

# Create summary table
summary = pd.DataFrame({
    "Metric": [
        "Total Sales",
        "Total Expenses",
        "Total Profit",
        "Profit Margin",
        "Total Orders",
        "Top Product",
        "Top Region"
    ],

    "Value": [
        total_sales,
        total_expenses,
        total_profit,
        profit_margin,
        total_orders,
        top_product,
        top_region
    ]
})


# Create Excel report
report_file = "Business_Report.xlsx"

with pd.ExcelWriter(
    report_file,
    engine="openpyxl"
) as writer:

    # Main business data
    df.to_excel(
        writer,
        sheet_name="Business Data",
        index=False
    )

    # KPI summary
    summary.to_excel(
        writer,
        sheet_name="Summary",
        index=False
    )


print()
print("================================")
print(" BUSINESS REPORT CREATED")
print("================================")
print("Report:", report_file)
print("================================")

from openpyxl import load_workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Font, Alignment
from openpyxl.chart import BarChart, Reference

# Open the generated workbook
wb = load_workbook("Business_Report.xlsx")

summary_sheet = wb["Summary"]
data_sheet = wb["Business Data"]


# ==============================
# FORMAT SUMMARY SHEET
# ==============================

# Make headers bold
for cell in summary_sheet[1]:
    cell.font = Font(bold=True)
    cell.alignment = Alignment(horizontal="center")


# Set column widths
summary_sheet.column_dimensions["A"].width = 25
summary_sheet.column_dimensions["B"].width = 25


# Format percentage
for row in range(2, summary_sheet.max_row + 1):

    if summary_sheet[f"A{row}"].value == "Profit Margin":
        summary_sheet[f"B{row}"].number_format = "0.00%"


# Format currency
for row in range(2, summary_sheet.max_row + 1):

    metric = summary_sheet[f"A{row}"].value

    if metric in [
        "Total Sales",
        "Total Expenses",
        "Total Profit"
    ]:
        summary_sheet[f"B{row}"].number_format = '₹#,##0'


# Freeze header
summary_sheet.freeze_panes = "A2"


# ==============================
# FORMAT BUSINESS DATA
# ==============================

data_sheet.freeze_panes = "A2"

for cell in data_sheet[1]:
    cell.font = Font(bold=True)
    cell.alignment = Alignment(horizontal="center")


# Format Sales and Expenses
for row in range(2, data_sheet.max_row + 1):

    data_sheet[f"D{row}"].number_format = '₹#,##0'
    data_sheet[f"E{row}"].number_format = '₹#,##0'


# Set column widths
data_sheet.column_dimensions["A"].width = 15
data_sheet.column_dimensions["B"].width = 18
data_sheet.column_dimensions["C"].width = 15
data_sheet.column_dimensions["D"].width = 15
data_sheet.column_dimensions["E"].width = 15


# ==============================
# ADD SALES CHART
# ==============================

chart = BarChart()

chart.title = "Sales by Product"

chart.y_axis.title = "Sales"

chart.x_axis.title = "Product"


# Create product sales data
product_sales = (
    df.groupby("Product")["Sales"]
    .sum()
    .reset_index()
)


# Put chart data into Summary sheet
start_row = summary_sheet.max_row + 3

summary_sheet[f"A{start_row}"] = "Product"
summary_sheet[f"B{start_row}"] = "Sales"


for i, row in product_sales.iterrows():

    summary_sheet[
        f"A{start_row + i + 1}"
    ] = row["Product"]

    summary_sheet[
        f"B{start_row + i + 1}"
    ] = row["Sales"]


# Chart data
data = Reference(
    summary_sheet,
    min_col=2,
    min_row=start_row,
    max_row=start_row + len(product_sales)
)

categories = Reference(
    summary_sheet,
    min_col=1,
    min_row=start_row + 1,
    max_row=start_row + len(product_sales)
)


chart.add_data(
    data,
    titles_from_data=True
)

chart.set_categories(categories)

chart.height = 7
chart.width = 12


summary_sheet.add_chart(
    chart,
    "D7"
)

# ==============================
# BUSINESS INSIGHTS
# ==============================

# Calculate product sales
product_sales = (
    df.groupby("Product")["Sales"]
    .sum()
)

# Calculate region sales
region_sales = (
    df.groupby("Region")["Sales"]
    .sum()
)

# Find top product
top_product = product_sales.idxmax()
top_product_sales = product_sales.max()

# Find top region
top_region = region_sales.idxmax()
top_region_sales = region_sales.max()


# Create insight messages

insight_1 = (
    f"{top_product} generated the highest sales "
    f"among all products with ₹{top_product_sales:,.0f}."
)

insight_2 = (
    f"{top_region} was the highest-sales region "
    f"with ₹{top_region_sales:,.0f} in sales."
)

insight_3 = (
    f"The business generated ₹{total_profit:,.0f} "
    f"profit from ₹{total_sales:,.0f} in sales."
)

insight_4 = (
    f"The overall profit margin was "
    f"{profit_margin:.2%}."
)


# ==============================
# KPI DASHBOARD CARDS
# ==============================

# Card 1 — Total Sales
summary_sheet.merge_cells("D2:E4")
summary_sheet["D2"] = f"TOTAL SALES\n₹{total_sales:,.0f}"

# Card 2 — Total Expenses
summary_sheet.merge_cells("F2:G4")
summary_sheet["F2"] = f"TOTAL EXPENSES\n₹{total_expenses:,.0f}"

# Card 3 — Total Profit
summary_sheet.merge_cells("H2:I4")
summary_sheet["H2"] = f"TOTAL PROFIT\n₹{total_profit:,.0f}"

# Card 4 — Profit Margin
summary_sheet.merge_cells("J2:K4")
summary_sheet["J2"] = f"PROFIT MARGIN\n{profit_margin:.2%}"


# Format KPI cards
card_cells = [
    "D2",
    "F2",
    "H2",
    "J2"
]

for cell in card_cells:

    summary_sheet[cell].font = Font(
        bold=True,
        size=14
    )

    summary_sheet[cell].alignment = Alignment(
        horizontal="center",
        vertical="center",
        wrap_text=True
    )


# Set row heights
summary_sheet.row_dimensions[2].height = 30
summary_sheet.row_dimensions[3].height = 30
summary_sheet.row_dimensions[4].height = 30


# Set column widths
for column in ["D", "E", "F", "G", "H", "I", "J", "K"]:

    summary_sheet.column_dimensions[
        column
    ].width = 15

# ==============================
# CONDITIONAL FORMATTING
# ==============================

# Sales color scale
data_sheet.conditional_formatting.add(
    f"D2:D{data_sheet.max_row}",
    ColorScaleRule(
        start_type="min",
        start_color="F8696B",
        mid_type="percentile",
        mid_value=50,
        mid_color="FFEB84",
        end_type="max",
        end_color="63BE7B"
    )
)


# Expenses color scale
data_sheet.conditional_formatting.add(
    f"E2:E{data_sheet.max_row}",
    ColorScaleRule(
        start_type="min",
        start_color="63BE7B",
        mid_type="percentile",
        mid_value=50,
        mid_color="FFEB84",
        end_type="max",
        end_color="F8696B"
    )
)


# ==============================
# WRITE BUSINESS INSIGHTS
# ==============================

insight_row = 20

summary_sheet[f"D{insight_row}"] = "BUSINESS INSIGHTS"

summary_sheet[f"D{insight_row + 1}"] = insight_1
summary_sheet[f"D{insight_row + 2}"] = insight_2
summary_sheet[f"D{insight_row + 3}"] = insight_3
summary_sheet[f"D{insight_row + 4}"] = insight_4


# Make the title bold
summary_sheet[f"D{insight_row}"].font = Font(
    bold=True,
    size=14
)


# Allow the text to wrap
for row in range(
    insight_row + 1,
    insight_row + 5
):
    summary_sheet[f"D{row}"].alignment = Alignment(
        wrap_text=True,
        vertical="top"
    )


# Make the insight area wider
summary_sheet.merge_cells("D21:K21")
summary_sheet.merge_cells("D22:K22")
summary_sheet.merge_cells("D23:K23")
summary_sheet.merge_cells("D24:K24")

# ==============================
# SAVE
# ==============================

wb.save("Business_Report.xlsx")

print("Professional dashboard formatting completed!")